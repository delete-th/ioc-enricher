# writes enrichment results to a docx file using python-docx

from pathlib import Path
from docx import Document


def write_report(results: dict) -> None:
    ioc = results["ioc"]
    ioc_type = results["type"]

    # sanitise ioc for use as a filename
    safe_name = ioc.replace("/", "_").replace(":", "_")
    output_path = Path(f"{safe_name}_report.docx")

    doc = Document()
    doc.add_heading("ioc enrichment report", level=1)
    doc.add_paragraph(f"ioc: {ioc}")
    doc.add_paragraph(f"type: {ioc_type}")

    # virustotal section
    doc.add_heading("virustotal", level=2)
    vt = results.get("virustotal")
    if vt and "error" not in vt:
        doc.add_paragraph(f"malicious: {vt['malicious']}")
        doc.add_paragraph(f"suspicious: {vt['suspicious']}")
        doc.add_paragraph(f"harmless: {vt['harmless']}")
        doc.add_paragraph(f"undetected: {vt['undetected']}")
    elif vt:
        doc.add_paragraph(f"error: {vt['error']}")
    else:
        doc.add_paragraph("not queried (no api key or unsupported ioc type)")

    # abuseipdb section — only populated for ip iocs
    doc.add_heading("abuseipdb", level=2)
    abuse = results.get("abuseipdb")
    if abuse and "error" not in abuse:
        doc.add_paragraph(f"confidence score: {abuse['abuse_confidence_score']}%")
        doc.add_paragraph(f"total reports: {abuse['total_reports']}")
        doc.add_paragraph(f"country: {abuse['country']}")
        doc.add_paragraph(f"isp: {abuse['isp']}")
        doc.add_paragraph(f"tor exit node: {abuse['is_tor']}")
    elif abuse:
        doc.add_paragraph(f"error: {abuse['error']}")
    else:
        doc.add_paragraph("not queried (ip iocs only)")

    # shodan section — only populated for ip iocs
    doc.add_heading("shodan", level=2)
    shodan = results.get("shodan")
    if shodan and "error" not in shodan:
        ports = ", ".join(map(str, shodan["ports"])) or "none"
        hostnames = ", ".join(shodan["hostnames"]) or "none"
        doc.add_paragraph(f"org: {shodan['org']}")
        doc.add_paragraph(f"country: {shodan['country']}")
        doc.add_paragraph(f"open ports: {ports}")
        doc.add_paragraph(f"hostnames: {hostnames}")
    elif shodan:
        doc.add_paragraph(f"error: {shodan['error']}")
    else:
        doc.add_paragraph("not queried (ip iocs only)")

    doc.save(output_path)
    print(f"[*] docx report saved to {output_path}")
